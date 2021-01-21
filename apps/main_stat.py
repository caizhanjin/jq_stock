"""
数据分析统计脚本
"""
from datetime import datetime
import time

from apps.clean_data import update_stock_info_to_database
from apps.orm_sqlite import JqBarData, JqStockInfo, StatDailyData
from peewee import JOIN
import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from jqdatasdk import *
from utils.utility import load_json, get_trade_days

from apps.jq_dataManage import save_stocks_info, save_history_data2, save_industry_info


def df_set_column_width(_writer, _sheet_name, df, is_filter=False):
    """宽度自适应"""
    worksheet = _writer.sheets[_sheet_name]
    if is_filter:
        (max_row, max_col) = df.shape
        worksheet.autofilter(0, 0, max_row, max_col - 1)

    worksheet.set_column(0, 0, 16)
    for i, col in enumerate(df.columns):
        collen = df[col].apply(lambda x: len(str(x).encode())).max()
        collen = max(len(col.encode()), collen)
        worksheet.set_column(i+1, i+1, min(collen + 3, 18))


def main(_start_date, _end_date=None):
    """
    统计脚本（主）
    """
    print(f"开始统计从 {_start_date} 到 {_end_date} 数据...")
    start_time = time.time()
    if _end_date is None:
        _end_date = datetime.now().strftime("%Y-%m-%d")

    old_stocks = load_json("../data/stocks_info.json")
    old_stocks_count = len(list(old_stocks.keys()))
    stocks = get_all_securities(['stock'])
    stocks_count = stocks.shape[0]
    if old_stocks_count != stocks_count:
        print(f"本地和聚宽股票数有差异，开始同步远端股票详情...")
        save_stocks_info(stocks)
        save_industry_info()
        update_stock_info_to_database()

    _, trade_days = get_trade_days(_start_date, _end_date)
    _bar_list = (
        JqBarData.select(JqBarData.datetime)
        .where((JqBarData.interval == "1d") & (JqBarData.index == "000001.XSHE"))
    )
    had_days = list(set([i.datetime.strftime("%Y-%m-%d") for i in _bar_list]))
    miss_days = [i for i in trade_days if i not in had_days]
    if len(miss_days) > 0:
        print(f"统计范围有缺失日期，开始同步缺失数据...")
        print(f"缺失日期：{str(miss_days)}")
        save_history_data2(miss_days[0], miss_days[-1])

    bar_list = (
        JqBarData.select(JqBarData, JqStockInfo)
        .join(
            JqStockInfo,
            on=(JqBarData.index == JqStockInfo.index),
            join_type=JOIN.LEFT_OUTER
        )
        .where((JqBarData.datetime in trade_days) & (JqStockInfo.type == "stock") & (JqBarData.interval == "1d"))
        .order_by(JqBarData.money.desc())
    )
    bar_dict = []

    for bar in bar_list:
        # 涨幅
        if bar.close and bar.pre_close:
            change_percent = round((bar.close - bar.pre_close) / bar.pre_close * 100, 2)
        else:
            change_percent = 0

        bar_dict.append({
            "code": bar.index,
            "display_name": "" if not bar.jqstockinfo else bar.jqstockinfo.display_name,
            "change_percent": change_percent,
            "money": 0 if not bar.money else bar.money,
            "datetime": bar.datetime,
            # "interval": bar.interval,

            "open": bar.open,
            "close": bar.close,
            "low": bar.low,
            "high": bar.high,
            "volume": bar.volume,
            "pre_close": bar.pre_close,

            "zjw_name": "" if not bar.jqstockinfo else bar.jqstockinfo.zjw_name,
            "sw_l1_name": "" if not bar.jqstockinfo else bar.jqstockinfo.sw_l1_name,
            "sw_l2_name": "" if not bar.jqstockinfo else bar.jqstockinfo.sw_l2_name,
            "sw_l3_name": "" if not bar.jqstockinfo else bar.jqstockinfo.sw_l3_name,

            "is_etf50": "否" if not bar.jqstockinfo else ("是" if bar.jqstockinfo.is_etf50 == 1 else "否"),
            "is_if300": "否" if not bar.jqstockinfo else ("是" if bar.jqstockinfo.is_if300 == 1 else "否"),
            "is_csi500": "否" if not bar.jqstockinfo else ("是" if bar.jqstockinfo.is_csi500 == 1 else "否"),
            "is_science": "否" if not bar.jqstockinfo else ("是" if bar.jqstockinfo.is_science == 1 else "否"),
        })

    all_df = pd.DataFrame.from_dict(bar_dict, orient="columns")
    all_df.sort_values(by=["money"], ascending=[False], inplace=True)

    volume_stat_list = []
    for trade_day in trade_days:
        _df = all_df[(all_df["datetime"] == trade_day + " 00:00:00")]

        df_1 = _df[(_df['money'] >= 100000000.0)]
        df_2 = _df[(_df['money'] >= 200000000.0)]
        df_20 = _df[_df["money"] >= 2000000000.0]

        volume_stat_list.append({
            "datetime": trade_day,

            "science_money": _df[(_df['is_science'] == '是')]['money'].sum(),
            "etf50_money": _df[(_df['is_etf50'] == '是')]['money'].sum(),
            "if300_money": _df[(_df['is_if300'] == '是')]['money'].sum(),
            "csi500_money": _df[(_df['is_csi500'] == '是')]['money'].sum(),

            "front10_money": _df[:10]['money'].sum(),
            "front15_money": _df[:15]['money'].sum(),
            "front20_money": _df[:20]['money'].sum(),

            "大于1亿数": df_1.shape[0],
            "大于2亿数": df_2.shape[0],
            "大于20亿数": df_20.shape[0],

            "大于1亿上涨数": df_1[(df_1['change_percent'] > 0)].shape[0],
            "大于1亿下跌数": -df_1[(df_1['change_percent'] < 0)].shape[0],
            "大于1亿持平数": df_1[(df_1['change_percent'] == 0)].shape[0],

            "大于2亿上涨数": df_2[(df_2['change_percent'] > 0)].shape[0],
            "大于2亿下跌数": -df_2[(df_2['change_percent'] < 0)].shape[0],
            "大于2亿持平数": df_2[(df_2['change_percent'] == 0)].shape[0],

            "大于20亿上涨数": df_20[(df_20['change_percent'] > 0)].shape[0],
            "大于20亿下跌数": -df_20[(df_20['change_percent'] < 0)].shape[0],
            "大于20亿持平数": df_20[(df_20['change_percent'] == 0)].shape[0],
        })

    order = [
        "code", "display_name", "change_percent", "money", "datetime",
        "open", "close", "low", "high", "volume", "pre_close",
        "zjw_name", "sw_l1_name", "sw_l2_name", "sw_l3_name",
        "is_etf50", "is_if300", "is_csi500", "is_science",
    ]
    stat_date = trade_days[-1]
    df = all_df[(all_df["datetime"] == stat_date + " 00:00:00")]
    df = df[order]
    df = df.set_index(["code"])
    df = df.drop(['datetime'], axis=1)
    df_original = df.copy(deep=True)
    df = df.sort_values(["sw_l1_name", "sw_l2_name", "sw_l3_name"], ascending=[1, 1, 1])

    df_1 = df[(df['money'] >= 100000000.0)]
    df_2 = df[(df['money'] >= 200000000.0)]
    df_20 = df[df["money"] >= 2000000000.0]

    # excel_writer初始化
    excel_name = f"{stat_date}_" \
                 f"{df_20.shape[0]}_" \
                 f"{df_2.shape[0]}_" \
                 f"{df_1.shape[0]}_" \
                 f"{int(df[(df['is_etf50'] == '是')]['money'].sum())}_" \
                 f"{int(df[(df['is_if300'] == '是')]['money'].sum())}_" \
                 f"{int(df[(df['is_csi500'] == '是')]['money'].sum())}_" \
                 f"{int(df[(df['is_science'] == '是')]['money'].sum())}" \
                 f".xlsx"
    excel_name_path = os.path.join("output_data", excel_name)
    excel_writer = pd.ExcelWriter(excel_name_path, engine='xlsxwriter')
    workbook = excel_writer.book
    header_format = workbook.add_format({
        'bold': True,
        'text_wrap': True,
        'align': 'center',
        'border': 1
    })
    title_format = workbook.add_format({
        'bold': True,
        'text_wrap': True,
        'align': 'center',
        'bg_color': '#92d050',
        'border': 1
    })

    # 成交前100、200涨跌家数和区间统计
    change_overall_dict = {}
    change_info_dict = {}
    change_info_array = ["前100", "前200", "大于1亿", "大于2亿", "大于20亿"]
    writer_array = ["大于1亿", "大于2亿", "大于20亿",
                    "100亿以上", "80(含)~100亿", "60(含)~80亿", "50(含)~60亿", "40(含)~50亿", "30(含)~40亿", "20(含)~30亿"]
    df_dict = {
        "前100": df_original[:100],
        "前200": df_original[:200],
        "大于1亿": df_1,
        "大于2亿": df_2,
        "大于20亿": df_20,
        "100亿以上": df[(df["money"] >= 10000000000.0)],
        "80(含)~100亿": df[(df["money"] >= 8000000000.0) & (df["money"] < 10000000000.0)],
        "60(含)~80亿": df[(df["money"] >= 6000000000.0) & (df["money"] < 8000000000.0)],
        "50(含)~60亿": df[(df["money"] >= 5000000000.0) & (df["money"] < 6000000000.0)],
        "40(含)~50亿": df[(df["money"] >= 4000000000.0) & (df["money"] < 5000000000.0)],
        "30(含)~40亿": df[(df["money"] >= 3000000000.0) & (df["money"] < 4000000000.0)],
        "20(含)~30亿": df[(df["money"] >= 2000000000.0) & (df["money"] < 3000000000.0)],
    }

    for key, item in df_dict.items():
        change_df = df_dict[key]
        change_overall_dict.setdefault(key, {})
        change_overall_dict[key] = {
            "stat_type": key,
            "count": change_df.shape[0],
            "up": change_df[(change_df["change_percent"] > 0)].shape[0],
            "down": -change_df[(change_df["change_percent"] < 0)].shape[0],
            "stay": change_df[(change_df["change_percent"] == 0)].shape[0],
        }
        # print(f"\n{key}涨跌家数和区间统计>>\n"
        #       f"总览：{change_overall_dict[key]}")

        if key in change_info_array:
            change_info_dict.setdefault(key, {})
            change_info_dict[key] = {
                "涨停": change_df[(change_df["change_percent"] >= 9.9)].shape[0],
                "7~9.9%": change_df[(change_df["change_percent"] >= 7) & (change_df["change_percent"] < 9.9)].shape[0],
                "5~7%": change_df[(change_df["change_percent"] >= 5) & (change_df["change_percent"] < 7)].shape[0],
                "2~5%": change_df[(change_df["change_percent"] >= 2) & (change_df["change_percent"] < 5)].shape[0],
                "0~2%": change_df[(change_df["change_percent"] > 0) & (change_df["change_percent"] < 2)].shape[0],
                "平": change_df[(change_df["change_percent"] == 0)].shape[0],
                "-0~2%": change_df[(change_df["change_percent"] > -2) & (change_df["change_percent"] < 0)].shape[0],
                "-2~5%": change_df[(change_df["change_percent"] > -5) & (change_df["change_percent"] <= -2)].shape[0],
                "-5~7%": change_df[(change_df["change_percent"] > -7) & (change_df["change_percent"] <= -5)].shape[0],
                "-7~9.9%": change_df[(change_df["change_percent"] > -9.9) & (change_df["change_percent"] <= -7)].shape[0],
                "跌停": change_df[(change_df["change_percent"] <= -9.9)].shape[0],
            }
            # print(f"详情：{change_info_dict[key]}")

    order = ["stat_type", "count", "up", "down", "stay"]
    df_change = pd.DataFrame.from_dict(list(change_overall_dict.values()), orient="columns")
    df_change = df_change[order]
    df_change.set_index(["stat_type"], inplace=True)
    df_change.to_excel(excel_writer, stat_date)
    df_set_column_width(excel_writer, stat_date, df_change)

    # daily_data数据查询和导出
    order = [
        "datetime",
        "science_money", "etf50_money", "if300_money", "csi500_money",
        "front10_money", "front15_money", "front20_money",
        "大于1亿数", "大于2亿数", "大于20亿数",
        "大于1亿上涨数", "大于1亿下跌数", "大于1亿持平数",
        "大于2亿上涨数", "大于2亿下跌数", "大于2亿持平数",
        "大于20亿上涨数", "大于20亿下跌数", "大于20亿持平数",
    ]
    df_daily = pd.DataFrame.from_dict(volume_stat_list, orient="columns")
    df_daily = df_daily[order]
    df_daily.set_index(["datetime"], inplace=True)
    df_daily.to_excel(excel_writer, "daily_data")

    # 画图
    worksheet = excel_writer.sheets[stat_date]
    column_chart = workbook.add_chart({'type': 'column', 'subtype': 'stacked'})
    column_chart.add_series({
        'name': f'上涨',
        'categories': f'={stat_date}!$A$2:$A$13',
        'values': f'={stat_date}!$C$2:$C$13',
        'fill': {'color': '#c00000'},
    })
    column_chart.add_series({
        'name': f'下跌',
        'categories': f'={stat_date}!$A$2:$A$13',
        'values': f'={stat_date}!$D$2:$D$13',
        'fill': {'color': '#00b050'},
    })
    line_chart = workbook.add_chart({'type': 'line'})
    line_chart.add_series({
        'name': '总数',
        'categories': f'={stat_date}!$A$2:$A$13',
        'values': f'={stat_date}!$B$2:$B$13',
        'data_labels': {'value': True},
    })

    column_chart.combine(line_chart)
    column_chart.set_title({'name': '区间统计总曲线'})
    column_chart.set_size({'width': 900, 'height': 400})
    worksheet.insert_chart(f'G1', column_chart)

    header = ["区间", "上涨", "持平", "下跌"]
    base_place = 25
    for key, item in change_info_dict.items():
        worksheet.merge_range(f'A{base_place}:D{base_place}', "区间细分统计——" + key, title_format)
        worksheet.write_row(f'A{base_place+1}', header, header_format)
        worksheet.write_column(f'A{base_place+2}', list(item.keys()), header_format)
        worksheet.write_column(f'B{base_place+2}', [
            item["涨停"], item["7~9.9%"], item["5~7%"], item["2~5%"], item["0~2%"],
            '-',
            '-', '-', '-', '-', '-'
        ])
        worksheet.write_column(f'C{base_place+2}', [
            '-', '-', '-', '-', '-',
            item["平"],
            '-', '-', '-', '-', '-'
        ])
        worksheet.write_column(f'D{base_place+2}', [
            '-', '-', '-', '-', '-',
            '-',
            item["-0~2%"], item["-2~5%"], item["-5~7%"], item["-7~9.9%"], item["跌停"],
        ])

        chart = workbook.add_chart({'type': 'column', 'subtype': 'stacked'})
        chart.add_series({
            'name': f'上涨',
            'categories': f'={stat_date}!$A${base_place+2}:$A${base_place+12}',
            'values': f'={stat_date}!$B${base_place+2}:$B${base_place+12}',
            'fill': {'color': '#c00000'},
        })
        chart.add_series({
            'name': f'持平',
            'categories': f'={stat_date}!$A${base_place+2}:$A${base_place+12}',
            'values': f'={stat_date}!$C${base_place+2}:$C${base_place+12}',
            'fill': {'color': '#808080'},
        })
        chart.add_series({
            'name': f'下跌',
            'categories': f'={stat_date}!$A${base_place+2}:$A${base_place+12}',
            'values': f'={stat_date}!$D${base_place+2}:$D${base_place+12}',
            'fill': {'color': '#00b050'},
        })

        chart.set_size({'width': 700, 'height': 300})
        worksheet.insert_chart(f'G{base_place}', chart)

        base_place += 18

    worksheet = excel_writer.sheets["daily_data"]
    trade_days_len = len(trade_days)
    base_place = trade_days_len + 5

    column_chart = workbook.add_chart({'type': 'column', 'subtype': 'stacked'})
    column_chart.add_series({
        'name': f'上涨数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$L$2:$L${trade_days_len+1}',
        'fill': {'color': '#c00000'},
    })
    column_chart.add_series({
        'name': f'下跌数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$M$2:$M${trade_days_len+1}',
        'fill': {'color': '#00b050'},
    })
    column_chart.add_series({
        'name': f'持平数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$N$2:$N${trade_days_len+1}',
        'fill': {'color': '#808080'},
    })
    line_chart = workbook.add_chart({'type': 'line'})
    line_chart.add_series({
        'name': '总数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$I$2:$I${trade_days_len+1}',
        'data_labels': {'value': True},
    })
    column_chart.combine(line_chart)
    column_chart.set_title({'name': '大于1亿统计总曲线'})
    column_chart.set_size({'width': 1100, 'height': 460})
    worksheet.insert_chart(f'A{base_place}', column_chart)

    base_place += 25
    column_chart = workbook.add_chart({'type': 'column', 'subtype': 'stacked'})
    column_chart.add_series({
        'name': f'上涨数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$O$2:$O${trade_days_len+1}',
        'fill': {'color': '#c00000'},
    })
    column_chart.add_series({
        'name': f'下跌数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$P$2:$P${trade_days_len+1}',
        'fill': {'color': '#00b050'},
    })
    column_chart.add_series({
        'name': f'持平数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$Q$2:$Q${trade_days_len+1}',
        'fill': {'color': '#808080'},
    })
    line_chart = workbook.add_chart({'type': 'line'})
    line_chart.add_series({
        'name': '总数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$J$2:$J${trade_days_len+1}',
        'data_labels': {'value': True},
    })
    column_chart.combine(line_chart)
    column_chart.set_title({'name': '大于2亿统计总曲线'})
    column_chart.set_size({'width': 1100, 'height': 460})
    worksheet.insert_chart(f'A{base_place}', column_chart)

    base_place += 25
    column_chart = workbook.add_chart({'type': 'column', 'subtype': 'stacked'})
    column_chart.add_series({
        'name': f'上涨数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$R$2:$R${trade_days_len+1}',
        'fill': {'color': '#c00000'},
    })
    column_chart.add_series({
        'name': f'下跌数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$S$2:$S${trade_days_len+1}',
        'fill': {'color': '#00b050'},
    })
    column_chart.add_series({
        'name': f'持平数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$T$2:$T${trade_days_len+1}',
        'fill': {'color': '#808080'},
    })
    line_chart = workbook.add_chart({'type': 'line'})
    line_chart.add_series({
        'name': '总数',
        'categories': f'=daily_data!$A$2:$A${trade_days_len+1}',
        'values': f'=daily_data!$K$2:$K${trade_days_len+1}',
        'data_labels': {'value': True},
    })
    column_chart.combine(line_chart)
    column_chart.set_title({'name': '大于20亿统计总曲线'})
    column_chart.set_size({'width': 1100, 'height': 460})
    worksheet.insert_chart(f'A{base_place}', column_chart)

    df_set_column_width(excel_writer, "daily_data", df_daily)
    df_original.to_excel(excel_writer, "All")  # 所有数据写入
    df_set_column_width(excel_writer, "All", df_original, True)

    for key, item in df_dict.items():
        if key in writer_array:
            change_df = df_dict[key]
            change_df.to_excel(excel_writer, key)
            df_set_column_width(excel_writer, key, change_df, True)

    excel_writer.save()

    print(f"\n完成统计，累计耗时 {round(time.time() - start_time, 2)} s...")
    pass


if __name__ == '__main__':
    """统计汇总，只需跑一次该脚本即可"""
    start_date = "2021-01-01"
    end_date = "2021-01-10"
    main(start_date, end_date)

    pass
